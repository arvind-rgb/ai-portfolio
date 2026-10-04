#!/usr/bin/env python3
"""
update_master_xlsx.py — Stage 5: roll campaign results into campaigns-master.xlsx.

Run by GitHub Actions after fetch_stats.py. Creates the workbook if it doesn't exist.

Usage:
  python3 update_master_xlsx.py
  python3 update_master_xlsx.py --xlsx path/to/custom.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent.parent
EMAIL_MARKETING = PROJECT_ROOT / "email-marketing"
DEFAULT_XLSX = EMAIL_MARKETING / "campaigns-master.xlsx"

# Column order for the master sheet
COLUMNS = [
    "Campaign Slug",
    "Campaign Name",
    "Status",
    "From",
    "Audience Count",
    "Start Date",
    "End Date",
    "Send Time IST",
    "Total Sent",
    "Delivered",
    "Delivery Rate %",
    "Opened",
    "Open Rate %",
    "Clicked",
    "Click Rate %",
    "Bounced",
    "Complained",
    "Last Updated",
]

BHOOKLE_ORANGE = "FF6B00"
HEADER_BG = "FF6B00"
HEADER_FG = "FFFFFF"
ALT_ROW_BG = "FFF3E0"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Update campaigns-master.xlsx with latest results.")
    p.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="Path to master workbook")
    return p.parse_args()


def load_campaign_data(campaign_dir: Path) -> dict | None:
    config_file = campaign_dir / "config.json"
    results_file = campaign_dir / "results.json"

    config = json.loads(config_file.read_text()) if config_file.is_file() else {}
    results = json.loads(results_file.read_text()) if results_file.is_file() else {}

    plan_file = campaign_dir / "send-plan.json"
    plan = json.loads(plan_file.read_text()) if plan_file.is_file() else {}

    if not config and not results:
        return None

    return {
        "Campaign Slug": campaign_dir.name,
        "Campaign Name": config.get("name", campaign_dir.name),
        "Status": config.get("status", "Unknown"),
        "From": config.get("from_email", "team@example.com"),
        "Audience Count": config.get("audience_count", plan.get("total", 0)),
        "Start Date": plan.get("start_date", ""),
        "End Date": plan.get("end_date", ""),
        "Send Time IST": config.get("send_time_ist", ""),
        "Total Sent": results.get("total_sent", 0),
        "Delivered": results.get("delivered", 0),
        "Delivery Rate %": results.get("delivery_rate", 0.0),
        "Opened": results.get("opened", 0),
        "Open Rate %": results.get("open_rate", 0.0),
        "Clicked": results.get("clicked", 0),
        "Click Rate %": results.get("click_rate", 0.0),
        "Bounced": results.get("bounced", 0),
        "Complained": results.get("complained", 0),
        "Last Updated": results.get("last_updated", ""),
    }


def apply_header_style(ws, row_idx: int, num_cols: int):
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        fill = PatternFill(fill_type="solid", fgColor=HEADER_BG)
        font = Font(bold=True, color=HEADER_FG)
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = 18
    except ImportError:
        pass


def apply_alt_row_style(ws, row_idx: int, num_cols: int):
    if row_idx % 2 == 0:
        try:
            from openpyxl.styles import PatternFill
            fill = PatternFill(fill_type="solid", fgColor=ALT_ROW_BG)
            for col in range(1, num_cols + 1):
                ws.cell(row=row_idx, column=col).fill = fill
        except ImportError:
            pass


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        sys.exit("error: openpyxl not installed. Run: pip install openpyxl")

    args = parse_args()
    xlsx_path = Path(args.xlsx)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)

    campaigns_dir = EMAIL_MARKETING / "campaigns"
    campaign_dirs = sorted(campaigns_dir.iterdir()) if campaigns_dir.is_dir() else []

    rows = []
    for d in campaign_dirs:
        if d.is_dir():
            data = load_campaign_data(d)
            if data:
                rows.append(data)

    if not rows:
        print("No campaign data found.")
        return 0

    # Load or create workbook
    if xlsx_path.is_file():
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        ws.delete_rows(1, ws.max_row)  # clear and rewrite
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Campaigns"

    # Header row
    ws.append(COLUMNS)
    apply_header_style(ws, 1, len(COLUMNS))

    # Data rows
    for i, row in enumerate(rows, start=2):
        ws.append([row.get(col, "") for col in COLUMNS])
        apply_alt_row_style(ws, i, len(COLUMNS))

    # Freeze header
    ws.freeze_panes = "A2"

    wb.save(xlsx_path)
    print(f"Updated: {xlsx_path} ({len(rows)} campaigns)")

    for row in rows:
        print(f"  {row['Campaign Slug']:<30} open_rate={row['Open Rate %']}%  "
              f"click_rate={row['Click Rate %']}%  status={row['Status']}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())

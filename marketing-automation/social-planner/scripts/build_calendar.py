#!/usr/bin/env python3
"""
build_calendar.py — Bhookle Social Planner Excel builder
Usage: python3 build_calendar.py <content_json_path> <output_xlsx_path>
Reads the JSON schema defined in references/excel-spec.md and writes a
colour-coded, branded Excel workbook with Instagram_Facebook and LinkedIn tabs.
"""

import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import DoughnutChart, Reference
    from openpyxl.chart.series import DataPoint
except ImportError:
    sys.exit("openpyxl not found. Run: pip3 install openpyxl")

# ── Brand colours ────────────────────────────────────────────────────────────
PURPLE      = "7030A0"
MUSTARD     = "FFC000"
GREEN_TINT  = "E2EFDA"
PURPLE_TINT = "E9D6F5"
GREY_FILL   = "F5F5F5"
WHITE       = "FFFFFF"
DARK_GREY   = "333333"

# ── Column definitions ───────────────────────────────────────────────────────
IG_COLUMNS = [
    ("PostID",            12),
    ("PostDate",          14),
    ("PostName",          28),
    ("PostFormat",        14),
    ("ContentCategory",   18),
    ("Generator",         16),
    ("PostHook",          40),
    ("PostDescription",   60),
    ("PostHashtag",       50),
    ("ExistingAssetLink", 30),
    ("AssetAction",       14),
    ("PostGenPrompt",     60),
    ("PostOutput",        20),
    ("PostStatus",        14),
    ("EvalTarget",        30),
    ("EvalHypothesis",    50),
    ("PostPerformance",   20),
]

LI_COLUMNS = [
    ("PostID",          12),
    ("PostDate",        14),
    ("PostName",        28),
    ("PostTheme",       24),
    ("PostFormat",      18),
    ("Generator",       16),
    ("PostHook",        40),
    ("PostDescription", 60),
    ("PostHashtag",     40),
    ("PostGenPrompt",   60),
    ("PostOutput",      20),
    ("PostStatus",      14),
    ("EvalTarget",      30),
    ("EvalHypothesis",  50),
    ("PostPerformance", 20),
]

# Columns left empty (grey fill)
EMPTY_COLS = {"PostOutput", "PostStatus", "PostPerformance"}

# Generator fill colours
HIGGSFIELD_FILL   = PatternFill("solid", fgColor="D6E4F7")   # light blue
NANOBANA_FILL     = PatternFill("solid", fgColor="FFF2CC")   # light yellow

def classify_generator(post: dict) -> str:
    """
    Higgsfield  — cinematic video work: original Reel generation OR existing video re-edit.
                  Signals: 'Higgsfield' in prompt, 'Re-edit brief' in prompt, or Reformat action on a Reel.
    Nano Banana — image / lightweight: Static images, Carousels, Stories, and
                  trending-format Reels (POV, poll, comparison) that don't need cinematic render.
                  Signals: 'Nano-banana' in prompt, or format is not Reel.
    """
    fmt    = (post.get("post_format") or "").strip()
    prompt = (post.get("post_gen_prompt") or "").lower()
    action = (post.get("asset_action") or "").strip()

    if fmt == "Reel":
        if "nano-banana" in prompt:
            return "Nano Banana"
        # Re-edits of existing video assets → Higgsfield
        if "re-edit brief" in prompt or action == "Reformat":
            return "Higgsfield"
        # Original Higgsfield cinematic renders
        if "higgsfield" in prompt:
            return "Higgsfield"
        return "Higgsfield"   # default for any Reel

    # Static, Carousel, Story, Long-form Text, Document Carousel → Nano Banana
    return "Nano Banana"

def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def header_font():
    return Font(bold=True, color=WHITE, size=10)

def body_font(bold=False):
    return Font(bold=bold, color=DARK_GREY, size=9)

def thin_border_left(hex_color, width="medium"):
    side = Side(style=width, color=hex_color)
    return Border(left=side)

def write_header_row(ws, columns):
    fill   = make_fill(PURPLE)
    font   = header_font()
    align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, (col_name, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill  = fill
        cell.font  = font
        cell.alignment = align

def row_fill_for_category(category):
    cat = (category or "").strip().lower()
    if "chef" in cat:
        return None, make_fill("FFFBE6"), True    # mustard left border, light bg, is_chef
    if "food" in cat:
        return None, make_fill(GREEN_TINT), False
    if "brand" in cat:
        return None, make_fill(PURPLE_TINT), False
    return None, None, False

def write_data_row(ws, row_num, columns, post):
    _, row_bg, is_chef = row_fill_for_category(post.get("content_category", ""))
    generator = classify_generator(post)

    for col_idx, (col_name, _) in enumerate(columns, start=1):
        # Map Excel column names to JSON keys
        key_map = {
            "postid":            "post_id",
            "postdate":          "post_date",
            "postname":          "post_name",
            "postformat":        "post_format",
            "contentcategory":   "content_category",
            "generator":         "__generator__",
            "posthook":          "post_hook",
            "postdescription":   "post_description",
            "posthashtag":       "post_hashtag",
            "existingassetlink": "existing_asset_link",
            "assetaction":       "asset_action",
            "postgenprompt":     "post_gen_prompt",
            "postoutput":        "post_output",
            "poststatus":        "post_status",
            "evaltarget":        "eval_target",
            "evalhypothesis":    "eval_hypothesis",
            "postperformance":   "post_performance",
            "posttheme":         "post_theme",
        }
        json_key = key_map.get(col_name.lower(), col_name.lower())
        value = generator if json_key == "__generator__" else post.get(json_key, "")

        cell = ws.cell(row=row_num, column=col_idx, value=value)
        cell.font = body_font()
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        if col_name in EMPTY_COLS:
            cell.fill = make_fill(GREY_FILL)
        elif col_name == "Generator":
            cell.fill = HIGGSFIELD_FILL if generator == "Higgsfield" else NANOBANA_FILL
            cell.font = Font(bold=True, color=DARK_GREY, size=9)
        elif row_bg:
            cell.fill = row_bg

        if col_name == "PostGenPrompt":
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Mustard left border for chef video rows
    if is_chef:
        ws.cell(row=row_num, column=1).border = Border(
            left=Side(style="medium", color=MUSTARD)
        )

def set_column_widths(ws, columns):
    for col_idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

def freeze_and_style(ws):
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

def build_ig_tab(wb, posts):
    ws = wb.create_sheet("Instagram_Facebook")
    write_header_row(ws, IG_COLUMNS)
    for i, post in enumerate(posts, start=2):
        write_data_row(ws, i, IG_COLUMNS, post)
        ws.row_dimensions[i].height = 60
    set_column_widths(ws, IG_COLUMNS)
    freeze_and_style(ws)
    return ws

def build_li_tab(wb, posts):
    ws = wb.create_sheet("LinkedIn")
    write_header_row(ws, LI_COLUMNS)
    for i, post in enumerate(posts, start=2):
        write_data_row(ws, i, LI_COLUMNS, post)
        ws.row_dimensions[i].height = 60
    set_column_widths(ws, LI_COLUMNS)
    freeze_and_style(ws)
    return ws

def build_summary_tab(wb, data):
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False

    summary = data.get("summary", {})
    month   = data.get("month", "")
    trend   = data.get("trend_scout_summary", "")

    # Title block
    ws.merge_cells("A1:H1")
    title = ws["A1"]
    title.value = f"Bhookle Content Calendar — {month}"
    title.font  = Font(bold=True, color=WHITE, size=16)
    title.fill  = make_fill(PURPLE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Trend scout summary
    ws.merge_cells("A3:H6")
    ts = ws["A3"]
    ts.value = f"Trend Scout Summary\n{trend}"
    ts.font  = Font(color=DARK_GREY, size=9)
    ts.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 60

    # KPI tiles
    kpis = [
        ("Total Posts",    summary.get("total_posts", 0)),
        ("IG/FB Posts",    summary.get("ig_fb_posts", 0)),
        ("LinkedIn Posts", summary.get("linkedin_posts", 0)),
        ("Reels",          summary.get("reel_count", 0)),
        ("Carousels",      summary.get("carousel_count", 0)),
        ("Static",         summary.get("static_count", 0)),
        ("Stories",        summary.get("story_count", 0)),
        ("Net New",        summary.get("net_new_posts", 0)),
        ("Reused Assets",  summary.get("reused_assets", 0)),
    ]
    row = 8
    ws.cell(row=row, column=1, value="Content Mix Overview").font = Font(bold=True, color=PURPLE, size=11)
    row += 1
    for label, val in kpis:
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = body_font(bold=True)
        val_cell = ws.cell(row=row, column=2, value=val)
        val_cell.font = Font(bold=True, color=DARK_GREY, size=14)
        val_cell.fill = make_fill(MUSTARD)
        val_cell.alignment = Alignment(horizontal="center")
        row += 1

    # Category breakdown
    row += 1
    ws.cell(row=row, column=1, value="Category Breakdown").font = Font(bold=True, color=PURPLE, size=11)
    row += 1
    cats = [
        ("Chef Videos",   summary.get("chef_video_count", 0)),
        ("Food Content",  summary.get("food_content_count", 0)),
        ("Brand/Other",   summary.get("brand_count", 0)),
    ]
    for label, val in cats:
        ws.cell(row=row, column=1, value=label).font = body_font()
        ws.cell(row=row, column=2, value=val).font   = body_font(bold=True)
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 14

    return ws

def main():
    if len(sys.argv) < 3:
        sys.exit("Usage: build_calendar.py <content_json> <output_xlsx>")

    json_path  = Path(sys.argv[1])
    xlsx_path  = Path(sys.argv[2])

    if not json_path.exists():
        sys.exit(f"Content JSON not found: {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    ig_posts = data.get("instagram_facebook", [])
    li_posts = data.get("linkedin", [])

    if not ig_posts:
        print("WARNING: No instagram_facebook posts found in JSON.")
    if not li_posts:
        print("WARNING: No linkedin posts found in JSON.")

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    build_summary_tab(wb, data)
    build_ig_tab(wb, ig_posts)
    build_li_tab(wb, li_posts)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"Saved: {xlsx_path}")
    print(f"  Instagram/Facebook rows: {len(ig_posts)}")
    print(f"  LinkedIn rows:           {len(li_posts)}")

if __name__ == "__main__":
    main()

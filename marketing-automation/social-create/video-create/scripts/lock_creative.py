#!/usr/bin/env python3
"""Lock a final reel: copy to assets/social/YYYY-MM/, set PostStatus + PostOutput.
Single source of truth for status changes.

Usage:
  lock_creative.py CALENDAR.xlsx IG-010 --src /tmp/vc/IG-010/final.mp4 \
      --output-dir assets/social/2026-06 [--sheet Instagram_Facebook]
"""
import argparse, os, shutil, sys
import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("calendar")
    ap.add_argument("post_id")
    ap.add_argument("--src", required=True)
    ap.add_argument("--output-dir", required=True)
    ap.add_argument("--sheet", default="Instagram_Facebook")
    a = ap.parse_args()

    if not os.path.exists(a.src):
        print(f"ERROR: source missing: {a.src}", file=sys.stderr); sys.exit(1)
    os.makedirs(a.output_dir, exist_ok=True)
    ext = os.path.splitext(a.src)[1] or ".mp4"
    dest = os.path.join(a.output_dir, f"{a.post_id}{ext}")
    shutil.copy2(a.src, dest)

    wb = openpyxl.load_workbook(a.calendar)
    ws = wb[a.sheet]
    hdr = {c.value: c.column for c in ws[1]}
    rel = os.path.relpath(dest, os.getcwd())
    found = False
    for row in ws.iter_rows(min_row=2):
        if row[0].value == a.post_id:
            row[hdr["PostStatus"] - 1].value = "Creative Locked"
            row[hdr["PostOutput"] - 1].value = rel
            found = True
            break
    if not found:
        print(f"ERROR: {a.post_id} not found in {a.sheet}", file=sys.stderr); sys.exit(1)
    wb.save(a.calendar)
    print(f"LOCKED {a.post_id} -> {rel}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
list_pending.py — Find the first N posts in a bhookle-social-planner Excel
where Generator == "Nano Banana" AND PostStatus is empty or "Yet to post".

Usage:
  python3 list_pending.py <xlsx_path> [--ig N] [--li N] [--json]

Defaults: --ig 7 --li 7. With --json, outputs structured data Claude can pipe
into the next step. Without --json, prints a human-readable summary.
"""
from __future__ import annotations
import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not installed. Run: pip install openpyxl")


def header_index(ws):
    return {
        (ws.cell(1, c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }


def is_pending(status_value) -> bool:
    if status_value is None:
        return True
    s = str(status_value).strip().lower()
    return s == "" or s.startswith("yet")


def collect_pending(ws, limit: int) -> list[dict]:
    cols = header_index(ws)
    required = ["PostID", "Generator", "PostStatus", "PostGenPrompt", "PostFormat"]
    for col in required:
        if col not in cols:
            sys.exit(f"missing column in sheet {ws.title!r}: {col}")

    out = []
    for r in range(2, ws.max_row + 1):
        post_id = ws.cell(r, cols["PostID"]).value
        if not post_id:
            continue
        generator = (ws.cell(r, cols["Generator"]).value or "").strip()
        status    = ws.cell(r, cols["PostStatus"]).value
        if generator.lower() != "nano banana":
            continue
        if not is_pending(status):
            continue

        out.append({
            "row":               r,
            "post_id":           post_id,
            "post_format":       ws.cell(r, cols["PostFormat"]).value,
            "post_name":         ws.cell(r, cols.get("PostName", 0)).value if cols.get("PostName") else "",
            "post_hook":         ws.cell(r, cols.get("PostHook", 0)).value if cols.get("PostHook") else "",
            "post_gen_prompt":   ws.cell(r, cols["PostGenPrompt"]).value or "",
            "status":            status or "",
        })
        if len(out) >= limit:
            break
    return out


def main():
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--ig", type=int, default=7)
    p.add_argument("--li", type=int, default=7)
    p.add_argument("--json", action="store_true")
    args = p.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        sys.exit(f"file not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {name.lower(): name for name in wb.sheetnames}

    ig_ws = wb[sheets["instagram_facebook"]] if "instagram_facebook" in sheets else None
    li_ws = wb[sheets["linkedin"]]           if "linkedin"           in sheets else None

    result = {
        "xlsx":      str(path.resolve()),
        "instagram": collect_pending(ig_ws, args.ig) if ig_ws else [],
        "linkedin":  collect_pending(li_ws, args.li) if li_ws else [],
    }

    if args.json:
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return

    print(f"\n📅 Source: {result['xlsx']}\n")
    print(f"━━━ Instagram/Facebook ({len(result['instagram'])} pending) ━━━")
    for p in result["instagram"]:
        print(f"  {p['post_id']:8} | {p['post_format']:10} | {p['post_name']}")
    print(f"\n━━━ LinkedIn ({len(result['linkedin'])} pending) ━━━")
    for p in result["linkedin"]:
        print(f"  {p['post_id']:8} | {p['post_format']:14} | {p['post_name']}")


if __name__ == "__main__":
    main()

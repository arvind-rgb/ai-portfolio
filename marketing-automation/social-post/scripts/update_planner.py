#!/usr/bin/env python3
"""
update_planner.py — Write PostStatus back to the content planner Excel.

Used by:
  - social_post_runner.py after scheduling Instagram posts
  - post_linkedin.py after publishing LinkedIn posts (from GitHub Actions)
"""

from pathlib import Path
from typing import Optional

import openpyxl


def find_planner_file(project_root: Path) -> Optional[Path]:
    """Find the most recent content planner xlsx in content-planner/."""
    planner_dir = project_root / "content-planner"
    if not planner_dir.exists():
        return None
    candidates = sorted(
        [f for f in planner_dir.glob("bhookle-content-calendar-*.xlsx")
         if not f.name.startswith("~$") and ".bak" not in f.name],
        reverse=True
    )
    return candidates[0] if candidates else None


def update_post_status(
    planner_path: Path,
    sheet_name: str,
    post_id: str,
    new_status: str,
    extra_fields: Optional[dict] = None,
):
    """
    Update the PostStatus cell for the given PostID in the Excel sheet.
    Optionally write additional columns via extra_fields {col_name: value}.
    """
    wb = openpyxl.load_workbook(planner_path)
    if sheet_name not in wb.sheetnames:
        print(f"  ⚠️  Sheet '{sheet_name}' not found in {planner_path.name}")
        wb.close()
        return

    ws = wb[sheet_name]

    # Find header row
    headers = {cell.value: cell.column for cell in ws[1] if cell.value}
    status_col = headers.get("PostStatus")
    id_col = headers.get("PostID")

    if not status_col or not id_col:
        print(f"  ⚠️  PostID or PostStatus column not found in sheet '{sheet_name}'")
        wb.close()
        return

    # Find the row for this post_id
    for row in ws.iter_rows(min_row=2):
        cell_id = row[id_col - 1]
        if str(cell_id.value).strip() == str(post_id).strip():
            # Update status
            row[status_col - 1].value = new_status
            # Update extra fields if provided
            if extra_fields:
                for col_name, value in extra_fields.items():
                    col_idx = headers.get(col_name)
                    if col_idx:
                        row[col_idx - 1].value = value
            break

    wb.save(planner_path)
    wb.close()

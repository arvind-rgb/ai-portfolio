#!/usr/bin/env python3
"""
post_linkedin.py — LinkedIn API poster for Bhookle.

LinkedIn REST API has no native scheduling endpoint. This script posts immediately.
For advance scheduling, see gen_github_actions.py which generates a GitHub Actions
workflow that runs this script at the right time.

Called by: social_post_runner.py (via GitHub Actions) OR directly for immediate posts.

Usage (from GitHub Actions):
  python3 post_linkedin.py --post-id LI-001 --planner content-planner/bhookle-content-calendar-2026-06.xlsx
"""

import argparse
import os
import sys
import time
from pathlib import Path

import requests

LI_API_BASE = "https://api.linkedin.com/v2"
LI_UPLOAD_BASE = "https://api.linkedin.com/v2/assets"

# Bhookle company page — auto-tagged on every personal post
BHOOKLE_COMPANY_URN = "urn:li:organization:<COMPANY_ID>"
BHOOKLE_TAG_TEXT = "@Bhookle"
BHOOKLE_TAG_SUFFIX = f"\n\nMore from {BHOOKLE_TAG_TEXT}"  # appended to every caption


def build_caption_with_bhookle_tag(caption: str) -> tuple[str, list]:
    """
    Append @Bhookle tag to caption and return (final_text, attributes).
    LinkedIn renders the tagged company as a clickable link and notifies the page.
    """
    final_text = f"{caption}{BHOOKLE_TAG_SUFFIX}"
    # Find the position of @Bhookle in the final string
    tag_start = final_text.rfind(BHOOKLE_TAG_TEXT)
    attributes = [{
        "start": tag_start,
        "length": len(BHOOKLE_TAG_TEXT),
        "value": {
            "com.linkedin.common.CompanyAttributedEntity": {
                "company": BHOOKLE_COMPANY_URN
            }
        }
    }]
    return final_text, attributes


def check_li_credentials() -> bool:
    token = os.getenv("LI_ACCESS_TOKEN", "")
    author_urn = os.getenv("LI_AUTHOR_URN", "")
    client_id = os.getenv("LINKEDIN_CLIENT_ID", "")
    client_secret = os.getenv("LINKEDIN_CLIENT_SECRET", "")
    if not token or not author_urn:
        print("  ✗ LI_ACCESS_TOKEN or LI_AUTHOR_URN not set in .env")
        return False

    try:
        # 1. Identity via OpenID Connect userinfo
        r = requests.get(
            "https://api.linkedin.com/v2/userinfo",
            headers={"Authorization": f"Bearer {token}"},
            timeout=10,
        )
        if r.status_code == 401:
            print("  ✗ LinkedIn token EXPIRED — re-run linkedin_auth.py (setup.md §2)")
            return False
        name = r.json().get("name", "") if r.status_code == 200 else author_urn

        # 2. Expiry check via introspectToken (requires client creds)
        days_left = None
        if client_id and client_secret:
            i = requests.post(
                "https://www.linkedin.com/oauth/v2/introspectToken",
                data={"token": token, "client_id": client_id, "client_secret": client_secret},
                headers={"Content-Type": "application/x-www-form-urlencoded"},
                timeout=10,
            ).json()
            if i.get("active") and "expires_at" in i:
                import time
                days_left = (i["expires_at"] - int(time.time())) // 86400
                if days_left < 0:
                    print(f"  ✗ LinkedIn token EXPIRED ({-days_left}d ago) — re-run linkedin_auth.py")
                    return False
                if days_left <= 7:
                    print(f"  ⚠ LinkedIn token expires in {days_left}d — re-auth soon (setup.md §2)")

        suffix = f" (token: {days_left}d left)" if days_left is not None else ""
        print(f"  ✓ LinkedIn connected as {name or author_urn}{suffix}")
        return True
    except Exception as e:
        print(f"  ✗ LinkedIn credential check failed: {e}")
        return False


def register_image_upload(token: str, author_urn: str) -> tuple[str, str]:
    """Register an image upload with LinkedIn. Returns (asset_urn, upload_url)."""
    r = requests.post(
        f"{LI_API_BASE}/assets?action=registerUpload",
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
            "X-Restli-Protocol-Version": "2.0.0",
        },
        json={
            "registerUploadRequest": {
                "owner": author_urn,
                "recipes": ["urn:li:digitalmediaRecipe:feedshare-image"],
                "serviceRelationships": [{
                    "identifier": "urn:li:userGeneratedContent",
                    "relationshipType": "OWNER"
                }]
            }
        },
        timeout=30,
    )
    data = r.json()
    if r.status_code not in (200, 201):
        raise RuntimeError(f"LinkedIn image registration failed: {data}")

    upload_url = (data["value"]["uploadMechanism"]
                  ["com.linkedin.digitalmedia.uploading.MediaUploadHttpRequest"]["uploadUrl"])
    asset_urn = data["value"]["asset"]
    return asset_urn, upload_url


def upload_image_to_linkedin(local_path: str, token: str, author_urn: str) -> str:
    """Upload image and return the asset URN."""
    asset_urn, upload_url = register_image_upload(token, author_urn)

    with open(local_path, "rb") as f:
        image_bytes = f.read()

    r = requests.put(
        upload_url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/octet-stream",
        },
        data=image_bytes,
        timeout=120,
    )
    if r.status_code not in (200, 201):
        raise RuntimeError(f"LinkedIn image upload failed: {r.text}")

    return asset_urn


def post_linkedin_now(
    post_id: str,
    caption: str,
    asset_local_path: str,
    post_format: str,
) -> dict:
    """
    Post immediately to LinkedIn.
    Returns: dict with success, post_urn, error.
    """
    token = os.getenv("LI_ACCESS_TOKEN", "")
    author_urn = os.getenv("LI_AUTHOR_URN", "")

    try:
        import re
        caption = re.sub(
            r'(https?://)?bhookleeats\.com(/\S*)?',
            'Download the Bhookle app',
            caption,
            flags=re.IGNORECASE
        )

        is_text_only = "text" in post_format.lower() and "image" not in post_format.lower()
        has_asset = asset_local_path and Path(asset_local_path).exists()

        # Auto-tag @Bhookle company page so every personal post links/notifies the brand
        final_text, attributes = build_caption_with_bhookle_tag(caption)

        # Build post body
        post_body = {
            "author": author_urn,
            "lifecycleState": "PUBLISHED",
            "specificContent": {
                "com.linkedin.ugc.ShareContent": {
                    "shareCommentary": {
                        "text": final_text,
                        "attributes": attributes,
                    },
                    "shareMediaCategory": "NONE",
                }
            },
            "visibility": {
                "com.linkedin.ugc.MemberNetworkVisibility": "PUBLIC"
            }
        }

        # Add image if available
        if has_asset and not is_text_only:
            ext = Path(asset_local_path).suffix.lower()
            if ext in (".jpg", ".jpeg", ".png", ".gif"):
                print(f"    Uploading image to LinkedIn ({Path(asset_local_path).name})...")
                asset_urn = upload_image_to_linkedin(asset_local_path, token, author_urn)
                post_body["specificContent"]["com.linkedin.ugc.ShareContent"]["shareMediaCategory"] = "IMAGE"
                post_body["specificContent"]["com.linkedin.ugc.ShareContent"]["media"] = [{
                    "status": "READY",
                    "description": {"text": ""},
                    "media": asset_urn,
                    "title": {"text": ""},
                }]

        r = requests.post(
            f"{LI_API_BASE}/ugcPosts",
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/json",
                "X-Restli-Protocol-Version": "2.0.0",
            },
            json=post_body,
            timeout=30,
        )

        if r.status_code not in (200, 201):
            return {"success": False, "error": f"HTTP {r.status_code}: {r.text}"}

        personal_urn = r.headers.get("x-restli-id", "")

        # Also cross-post to Bhookle company page (if token has w_organization_social).
        # This is the reliable way to get the post on the Bhookle feed — far more
        # effective than relying on a personal-post @mention.
        company_urn_env = os.getenv("LI_COMPANY_URN", "")
        company_post_urn = None
        company_error = None
        if company_urn_env:
            # Build a company-authored body with attribution to Arvind in the text
            company_text = f"{caption}\n\n— Posted by Arvind, Founder at Bhookle"
            company_body = {
                "author": company_urn_env,
                "lifecycleState": "PUBLISHED",
                "specificContent": {
                    "com.linkedin.ugc.ShareContent": {
                        "shareCommentary": {"text": company_text},
                        "shareMediaCategory": "NONE",
                    }
                },
                "visibility": {"com.linkedin.ugc.MemberNetworkVisibility": "PUBLIC"},
            }
            # Re-upload image under company author URN (LinkedIn requires the owner of
            # the registered asset to match the post author)
            if has_asset and not is_text_only:
                ext = Path(asset_local_path).suffix.lower()
                if ext in (".jpg", ".jpeg", ".png", ".gif"):
                    try:
                        print(f"    Cross-posting to Bhookle page — uploading image under company URN...")
                        c_asset_urn = upload_image_to_linkedin(asset_local_path, token, company_urn_env)
                        company_body["specificContent"]["com.linkedin.ugc.ShareContent"]["shareMediaCategory"] = "IMAGE"
                        company_body["specificContent"]["com.linkedin.ugc.ShareContent"]["media"] = [{
                            "status": "READY",
                            "description": {"text": ""},
                            "media": c_asset_urn,
                            "title": {"text": ""},
                        }]
                    except Exception as e:
                        company_error = f"company image upload failed: {e}"

            if company_error is None:
                cr = requests.post(
                    f"{LI_API_BASE}/ugcPosts",
                    headers={
                        "Authorization": f"Bearer {token}",
                        "Content-Type": "application/json",
                        "X-Restli-Protocol-Version": "2.0.0",
                    },
                    json=company_body,
                    timeout=30,
                )
                if cr.status_code in (200, 201):
                    company_post_urn = cr.headers.get("x-restli-id", "")
                    print(f"    ✓ Cross-posted to Bhookle page: {company_post_urn}")
                else:
                    company_error = f"HTTP {cr.status_code}: {cr.text[:300]}"
                    print(f"    ⚠ Cross-post to Bhookle page FAILED: {company_error}")
                    print(f"      (Personal post still succeeded. Likely cause: token lacks")
                    print(f"       w_organization_social scope — re-run linkedin_auth.py)")

        return {
            "success": True,
            "post_urn": personal_urn,
            "company_post_urn": company_post_urn,
            "company_error": company_error,
        }

    except Exception as e:
        return {"success": False, "error": str(e)}


def post_from_planner(post_id: str, planner_path: str):
    """Called by GitHub Actions — read post from planner and publish."""
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    from update_planner import update_post_status, find_planner_file
    from dotenv import load_dotenv
    import openpyxl
    import re

    load_dotenv()

    planner = Path(planner_path) if planner_path else find_planner_file(Path.cwd())
    if not planner or not planner.exists():
        print(f"❌ Planner not found: {planner_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(planner, read_only=True, data_only=True)
    ws = wb["LinkedIn"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    post_data = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        if str(d.get("PostID", "")).strip() == post_id:
            post_data = d
            break
    wb.close()

    if not post_data:
        print(f"❌ Post {post_id} not found in LinkedIn sheet")
        sys.exit(1)

    # IDEMPOTENCY — already posted? exit cleanly
    if str(post_data.get("PostStatus", "")).strip() == "Posted":
        print(f"⏭  {post_id} already Posted — skipping (idempotent)")
        return

    import re
    caption = str(post_data.get("PostDescription", "")).strip()
    hashtags = str(post_data.get("PostHashtag", "")).strip()
    if hashtags:
        caption = f"{caption}\n\n{hashtags}"
    caption = re.sub(
        r'(https?://)?bhookleeats\.com(/\S*)?',
        'Download the Bhookle app',
        caption,
        flags=re.IGNORECASE
    )

    # Normalize laptop-absolute PostOutput to repo-relative so it works on GH runner
    raw_asset = str(post_data.get("PostOutput") or "")
    marker = "/bhookle-marketing/"
    idx = raw_asset.find(marker)
    asset_path = raw_asset[idx + len(marker):] if idx >= 0 else raw_asset

    # JIT-fetch from Cloudinary if local file missing (covers GH Actions runner)
    if asset_path:
        sys.path.insert(0, str(Path(__file__).parent))
        from post_instagram import _ensure_asset_available
        try:
            for seg in [s.strip() for s in asset_path.split(",")]:
                _ensure_asset_available(seg)
        except FileNotFoundError as e:
            print(f"  ⚠ {e}")

    result = post_linkedin_now(
        post_id=post_id,
        caption=caption,
        asset_local_path=asset_path,
        post_format=str(post_data.get("PostFormat") or ""),
    )

    if result["success"]:
        # Append to posted-log FIRST (idempotency safety net), then update planner.
        try:
            from social_tick import append_to_posted_log
            append_to_posted_log(post_id, "linkedin", result.get("post_urn", ""))
            print(f"  → Appended {post_id} to posted-log.csv (idempotency)")
        except Exception as e:
            print(f"  ⚠ Could not append to posted-log: {e}")
        update_post_status(planner, "LinkedIn", post_id, "Posted")
        print(f"✅ {post_id} posted successfully. URN: {result.get('post_urn', '')}")
    else:
        update_post_status(planner, "LinkedIn", post_id, "Failed")
        print(f"❌ {post_id} failed: {result['error']}")
        sys.exit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--post-id", required=True)
    parser.add_argument("--planner", default="")
    args = parser.parse_args()
    post_from_planner(args.post_id, args.planner)

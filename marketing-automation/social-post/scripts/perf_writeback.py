#!/usr/bin/env python3
"""perf_writeback.py — measure what we can with current scopes, write it back.

Runs daily on GH Actions (also fine locally). For every Instagram post in the
planner that went live ≥72h ago and has no recorded performance yet:

  - likes + comments via Graph API (basic scopes — no insights needed)
  - "interactions @72h+" vs the row's numeric target → achievement %
  - DABBA keyword count in comments (Kids CTA signal; instagram_manage_comments)

Also appends one account-level row per run to social/growth-log.csv
(IG followers via /{ig-user}, FB page followers via /{page}) so the
before/after report has a daily curve, not two points.

Reach / saves / shares / views are NOT fetched — token lacks
instagram_manage_insights. When that scope lands, extend fetch_media().
LinkedIn rows are skipped (no analytics scopes; manual count only).
"""
import csv
import datetime as dt
import json
import os
import re
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PLANNER = ROOT / "content-planner" / "bhookle-content-calendar-2026-08.xlsx"
POSTED_LOG = ROOT / "assets" / "_posted-log.csv"
GROWTH_LOG = ROOT / "social" / "growth-log.csv"
HOURS_MIN = 72

TOK = os.environ.get("IG_ACCESS_TOKEN")
UID = os.environ.get("IG_USER_ID")
PAGE = os.environ.get("FB_PAGE_ID", "<FB_PAGE_ID>")


def _load_env():
    envf = ROOT / ".env"
    if not envf.exists():
        return
    for line in envf.read_text().splitlines():
        line = line.strip()
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def get(url):
    return json.load(urllib.request.urlopen(url, timeout=30))


def fetch_media(media_id):
    d = get(f"https://graph.facebook.com/v21.0/{media_id}"
            f"?fields=like_count,comments_count,permalink&access_token={TOK}")
    dabba = 0
    try:
        c = get(f"https://graph.facebook.com/v21.0/{media_id}/comments"
                f"?fields=text&limit=100&access_token={TOK}")
        dabba = sum(1 for x in c.get("data", []) if "dabba" in (x.get("text") or "").lower())
    except Exception:
        pass
    return d.get("like_count", 0), d.get("comments_count", 0), dabba


def growth_snapshot():
    row = {"date": dt.date.today().isoformat(), "ig_followers": "", "fb_followers": "",
           "li_followers_manual": ""}
    try:
        row["ig_followers"] = get(f"https://graph.facebook.com/v21.0/{UID}"
                                  f"?fields=followers_count&access_token={TOK}")["followers_count"]
    except Exception as e:
        print(f"  ig followers fetch failed: {e}")
    try:
        row["fb_followers"] = get(f"https://graph.facebook.com/v21.0/{PAGE}"
                                  f"?fields=followers_count&access_token={TOK}")["followers_count"]
    except Exception as e:
        print(f"  fb followers fetch failed: {e}")
    GROWTH_LOG.parent.mkdir(exist_ok=True)
    exists = GROWTH_LOG.exists()
    lines = GROWTH_LOG.read_text().splitlines() if exists else []
    if any(l.startswith(row["date"] + ",") for l in lines[1:]):
        print(f"  growth-log already has {row['date']}")
        return
    with open(GROWTH_LOG, "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(row))
        if not exists:
            w.writeheader()
        w.writerow(row)
    print(f"  growth-log +{row['date']}: IG {row['ig_followers']} | FB {row['fb_followers']}")


def main():
    _load_env()
    global TOK, UID
    TOK = TOK or os.environ.get("IG_ACCESS_TOKEN")
    UID = UID or os.environ.get("IG_USER_ID")
    if not TOK or not UID:
        sys.exit("IG_ACCESS_TOKEN / IG_USER_ID missing")

    posted = {}  # PostID -> (media_id, posted_at)
    if POSTED_LOG.exists():
        for r in csv.reader(open(POSTED_LOG)):
            if len(r) >= 4 and r[1] == "instagram":
                try:
                    posted[r[0]] = (r[2], dt.datetime.fromisoformat(r[3]))
                except ValueError:
                    continue  # header / seed rows

    import openpyxl
    wb = openpyxl.load_workbook(PLANNER)
    ws = wb["Instagram_Facebook"]
    h = [c.value for c in ws[1]]
    col = {k: h.index(k) + 1 for k in ("PostID", "PostPerformance", "EvalTarget")}
    now = dt.datetime.now(dt.timezone.utc)
    wrote = 0
    for r in range(2, ws.max_row + 1):
        pid = str(ws.cell(r, col["PostID"]).value or "").strip()
        if pid not in posted:
            continue
        if str(ws.cell(r, col["PostPerformance"]).value or "").strip():
            continue  # already measured — 72h snapshot is immutable
        media_id, posted_at = posted[pid]
        age_h = (now - posted_at).total_seconds() / 3600
        if age_h < HOURS_MIN:
            print(f"  {pid}: {age_h:.0f}h old — waiting for 72h mark")
            continue
        try:
            likes, comments, dabba = fetch_media(media_id)
        except Exception as e:
            print(f"  {pid}: fetch failed {e}")
            continue
        inter = likes + comments
        target = None
        m = re.search(r"Interactions:\s*(\d+)", str(ws.cell(r, col["EvalTarget"]).value or ""))
        if m:
            target = int(m.group(1))
        verdict = f" | vs target {target}: {inter / target * 100:.0f}%" if target else ""
        dtxt = f" | DABBA comments: {dabba}" if dabba else ""
        ws.cell(r, col["PostPerformance"]).value = (
            f"Likes {likes} | Comments {comments} | Interactions {inter} @{age_h:.0f}h"
            f"{verdict}{dtxt} | reach/saves n/a (no insights scope) | measured {now.date()}")
        print(f"  {pid}: {inter} interactions ({likes}L+{comments}C){verdict}{dtxt}")
        wrote += 1
    if wrote:
        wb.save(PLANNER)
    print(f"writeback: {wrote} row(s) recorded")
    growth_snapshot()


if __name__ == "__main__":
    main()
